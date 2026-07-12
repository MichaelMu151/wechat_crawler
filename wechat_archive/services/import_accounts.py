from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from wechat_archive.db import Database


def import_name_list(db: Database, xlsx_path: str | Path) -> dict[str, int]:
    """从 name_list.xlsx 导入 nickname + link。"""
    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"名单文件不存在: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"inserted": 0, "skipped": 0, "total": 0}

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    # 兼容 nickname/link 或 nickname/sample_url
    try:
        nick_i = header.index("nickname")
    except ValueError as e:
        raise ValueError("Excel 需要 nickname 列") from e

    link_i = None
    for cand in ("link", "sample_url", "url"):
        if cand in header:
            link_i = header.index(cand)
            break
    if link_i is None:
        raise ValueError("Excel 需要 link / sample_url / url 列")

    inserted = skipped = 0
    with db.connection() as conn:
        for row in rows[1:]:
            if not row or row[nick_i] is None:
                continue
            nickname = str(row[nick_i]).strip()
            if not nickname:
                continue
            link = ""
            if row[link_i] is not None:
                link = str(row[link_i]).strip()

            exists = conn.execute(
                "SELECT id FROM accounts WHERE nickname_input = ? AND IFNULL(sample_url,'') = ?",
                (nickname, link),
            ).fetchone()
            if exists:
                skipped += 1
                continue

            conn.execute(
                """
                INSERT INTO accounts (nickname_input, sample_url, resolve_status, list_status)
                VALUES (?, ?, 'pending', 'pending')
                """,
                (nickname, link or None),
            )
            inserted += 1

    return {"inserted": inserted, "skipped": skipped, "total": inserted + skipped}
